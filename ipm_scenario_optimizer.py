#!/usr/bin/env python3
"""IPM scenario optimiser V1.

Input workbook named tables:
  tblWorkers, tblScopeConfig, tblScenarioConfig

Dependencies: pandas numpy scipy openpyxl pyarrow
Usage: python ipm_scenario_optimizer.py --input inputs.xlsx --output ./output
"""
from __future__ import annotations

import argparse, hashlib, json, sys, time, uuid
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import linprog, minimize

RATINGS = [1, 2, 3, 4, 5]
VERSION = "1.0.0"


def fail(msg):
    raise ValueError(msg)


def as_bool(v):
    if isinstance(v, (bool, np.bool_)): return bool(v)
    s = str(v).strip().lower()
    if s in {"y", "yes", "true", "1", "active"}: return True
    if s in {"n", "no", "false", "0", "inactive"}: return False
    fail(f"Invalid boolean value: {v!r}")


def read_table(path: Path, name: str) -> pd.DataFrame:
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        for ws in wb.worksheets:
            if name in ws.tables:
                values = [[c.value for c in row] for row in ws[ws.tables[name].ref]]
                headers = [str(x).strip() if x is not None else "" for x in values[0]]
                if any(not x for x in headers) or len(headers) != len(set(headers)):
                    fail(f"{name}: blank or duplicate header")
                return pd.DataFrame(values[1:], columns=headers).dropna(how="all").reset_index(drop=True)
    finally:
        wb.close()
    fail(f"Named table {name!r} not found")


def need(df, table, cols):
    missing = [c for c in cols if c not in df.columns]
    if missing: fail(f"{table}: missing columns {missing}")


def prepare(workers, scopes, scenarios):
    need(workers, "tblWorkers", ["WorkerID","GroupExecutive","SubGroup","PerformanceRating","ProratedVRTarget","Eligible"])
    need(scopes, "tblScopeConfig", ["ScopeID","ScopeType","GroupExecutive","SubGroup","ParentScopeID","GEHoldbackPct","Active"])
    need(scenarios, "tblScenarioConfig", ["ScenarioID","ScenarioName","ScopeID","ShapeName","DifferentiationLevel","PerformanceRating","MinIPM","MaxIPM","ShapeScore","BaselineFlag","Active"])
    w, s, c = workers.copy(), scopes.copy(), scenarios.copy()
    for col in ["WorkerID","GroupExecutive","SubGroup"]: w[col] = w[col].fillna("").astype(str).str.strip()
    w["Eligible"] = w["Eligible"].map(as_bool)
    w["PerformanceRating"] = pd.to_numeric(w["PerformanceRating"]).astype(int)
    w["ProratedVRTarget"] = pd.to_numeric(w["ProratedVRTarget"]).astype(float)
    if w["WorkerID"].eq("").any() or w["WorkerID"].duplicated().any(): fail("WorkerID must be unique and non-blank")
    if not set(w["PerformanceRating"]).issubset(RATINGS): fail("PerformanceRating must be 1-5")
    if w["ProratedVRTarget"].lt(0).any(): fail("ProratedVRTarget cannot be negative")
    w["EligibleVRTarget"] = np.where(w["Eligible"], w["ProratedVRTarget"], 0.0)

    for col in ["ScopeID","ScopeType","GroupExecutive","SubGroup","ParentScopeID"]: s[col] = s[col].fillna("").astype(str).str.strip()
    s["Active"] = s["Active"].map(as_bool)
    s["ScopeType"] = s["ScopeType"].str.lower().map({"ge":"GE","subgroup":"SubGroup"})
    s["GEHoldbackPct"] = pd.to_numeric(s["GEHoldbackPct"], errors="coerce")
    if s["ScopeID"].eq("").any() or s["ScopeID"].duplicated().any(): fail("ScopeID must be unique and non-blank")
    active_s = s[s["Active"]].copy()
    if active_s.empty: fail("No active scopes")
    idx = active_s.set_index("ScopeID")
    for row in active_s.itertuples():
        if row.ScopeType == "GE":
            if pd.isna(row.GEHoldbackPct) or not 0 <= row.GEHoldbackPct <= 1: fail(f"{row.ScopeID}: invalid GEHoldbackPct")
        elif row.ScopeType == "SubGroup":
            if pd.notna(row.GEHoldbackPct): fail(f"{row.ScopeID}: subgroup holdback must be blank")
            if row.ParentScopeID not in idx.index or idx.loc[row.ParentScopeID,"ScopeType"] != "GE": fail(f"{row.ScopeID}: invalid parent GE")
        else: fail(f"{row.ScopeID}: ScopeType must be GE or SubGroup")

    for col in ["ScenarioID","ScenarioName","ScopeID","ShapeName","DifferentiationLevel"]: c[col] = c[col].fillna("").astype(str).str.strip()
    c["Active"] = c["Active"].map(as_bool); c["BaselineFlag"] = c["BaselineFlag"].map(as_bool)
    c["PerformanceRating"] = pd.to_numeric(c["PerformanceRating"]).astype(int)
    for col in ["MinIPM","MaxIPM","ShapeScore"]: c[col] = pd.to_numeric(c[col]).astype(float)
    active_c = c[c["Active"]].copy()
    if active_c.empty: fail("No active scenarios")
    if not set(active_c["ScopeID"]).issubset(set(active_s["ScopeID"])): fail("Scenario config references unknown/inactive ScopeID")
    for (sid, scope), g in active_c.groupby(["ScenarioID","ScopeID"]):
        g = g.sort_values("PerformanceRating")
        if g["PerformanceRating"].tolist() != RATINGS: fail(f"{sid}/{scope}: requires exactly ratings 1-5")
        if g["MinIPM"].gt(g["MaxIPM"]).any(): fail(f"{sid}/{scope}: MinIPM > MaxIPM")
        if np.any(np.diff(g["ShapeScore"]) < -1e-12): fail(f"{sid}/{scope}: ShapeScore must be non-decreasing")
        x = -np.inf
        for mn, mx in zip(g["MinIPM"], g["MaxIPM"]):
            x = max(x, mn)
            if x > mx + 1e-12: fail(f"{sid}/{scope}: min/max bounds cannot satisfy rating ordering")
    base = active_c[active_c["BaselineFlag"]].drop_duplicates(["ScenarioID","ScopeID"]).groupby("ScopeID")["ScenarioID"].nunique()
    if (base > 1).any(): fail("Only one baseline scenario is allowed per scope")
    return w, s, c


def in_scope(workers, scope):
    m = workers["GroupExecutive"].eq(scope.GroupExecutive)
    if scope.ScopeType == "SubGroup": m &= workers["SubGroup"].eq(scope.SubGroup)
    return workers[m].copy()


def holdback(scope, scope_idx):
    return float(scope.GEHoldbackPct if scope.ScopeType == "GE" else scope_idx.loc[scope.ParentScopeID,"GEHoldbackPct"])


def aggregate(workers):
    e = workers[workers["Eligible"]]
    g = e.groupby("PerformanceRating").agg(EligibleHC=("WorkerID","count"), EligibleVRTarget=("EligibleVRTarget","sum"))
    return pd.DataFrame({"PerformanceRating":RATINGS}).merge(g, left_on="PerformanceRating", right_index=True, how="left").fillna(0)


def solve(agg, cfg, budget, spend_tolerance):
    cfg = cfg.sort_values("PerformanceRating"); agg = agg.sort_values("PerformanceRating")
    vr = agg["EligibleVRTarget"].to_numpy(float); mn = cfg["MinIPM"].to_numpy(float); mx = cfg["MaxIPM"].to_numpy(float); shape = cfg["ShapeScore"].to_numpy(float)
    minimum = float(vr @ mn)
    if minimum > budget + 1e-8:
        return {"status":"INFEASIBLE_MINIMUM_SPEND","minimum":minimum,"shortfall":minimum-budget}
    order = []
    for i in range(4):
        row = np.zeros(5); row[i]=1; row[i+1]=-1; order.append(row)
    st1 = linprog(-vr, A_ub=np.vstack([vr, order]), b_ub=np.array([budget,0,0,0,0]), bounds=list(zip(mn,mx)), method="highs")
    if not st1.success: return {"status":"SOLVER_STAGE1_ERROR","message":st1.message,"minimum":minimum}
    max_spend = float(vr @ st1.x); floor = max(0, max_spend - spend_tolerance)
    if np.ptp(shape) > 1e-12:
        a0, k0 = np.linalg.lstsq(np.c_[np.ones(5),shape], st1.x, rcond=None)[0]; k0=max(0,float(k0)); a0=float(a0)
    else: a0=float(np.mean(st1.x)); k0=0.0
    def obj(z): return float(np.sum((z[:5]-(z[5]+z[6]*shape))**2))
    cons=[{"type":"ineq","fun":lambda z: budget-vr@z[:5]}, {"type":"ineq","fun":lambda z: vr@z[:5]-floor}]
    cons += [{"type":"ineq","fun":lambda z,i=i: z[i+1]-z[i]} for i in range(4)]
    st2 = minimize(obj, np.r_[st1.x,a0,k0], method="SLSQP", bounds=list(zip(mn,mx))+[(None,None),(0,None)], constraints=cons, options={"maxiter":1000,"ftol":1e-12})
    if not st2.success: return {"status":"SOLVER_STAGE2_ERROR","message":st2.message,"minimum":minimum,"max_spend":max_spend}
    x=st2.x[:5]; spend=float(vr@x); expected=st2.x[5]+st2.x[6]*shape
    rmse=float(np.sqrt(np.mean((x-expected)**2))); scale=max(float(mx.max()-mn.min()),1e-9)
    return {"status":"SUCCESS","ipm":dict(zip(RATINGS,map(float,x))),"minimum":minimum,"shortfall":0.0,"max_spend":max_spend,"spend":spend,"anchor":float(st2.x[5]),"curve_scale":float(st2.x[6]),"policy_fit":float(np.clip(100*(1-rmse/scale),0,100))}


def headroom(cfg, ipm):
    vals=[]
    for r in cfg.itertuples():
        width=r.MaxIPM-r.MinIPM
        vals.append(0 if width <= 1e-12 else min(ipm[r.PerformanceRating]-r.MinIPM,r.MaxIPM-ipm[r.PerformanceRating])/(width/2))
    return float(np.clip(np.mean(vals)*100,0,100))


def dominates(a,b,metrics):
    return all(a[m] >= b[m]-1e-10 for m in metrics) and any(a[m] > b[m]+1e-10 for m in metrics)


def pareto(ge):
    if ge.empty: return ge
    out=[]
    for (_, executive), g in ge.groupby(["BatchRunID","GroupExecutive"]):
        g=g.reset_index(drop=True); metrics=["BudgetUtilisationPct","PolicyFitScore","ConstraintHeadroomScore"]
        if g["StabilityScore"].notna().all(): metrics.append("StabilityScore")
        left=list(g.index); rank=1; ranks={}; by={i:None for i in left}
        while left:
            front=[]
            for i in left:
                dom=[j for j in left if j!=i and dominates(g.loc[j],g.loc[i],metrics)]
                if not dom: front.append(i)
                elif by[i] is None: by[i]=g.loc[dom[0],"ScenarioID"]
            if not front: front=left[:]
            for i in front: ranks[i]=rank
            left=[i for i in left if i not in front]; rank+=1
        for i,row in g.iterrows():
            d=row.to_dict(); d.update(ParetoRank=ranks[i],ParetoEfficientFlag=ranks[i]==1,DominatedFlag=ranks[i]>1,DominatedByScenarioID=by[i],ParetoMetricsUsed=",".join(metrics)); out.append(d)
    return pd.DataFrame(out)


def run_engine(workers, scopes, scenarios, spend_tolerance=1.0, tol=1e-6):
    w,s,c=prepare(workers,scopes,scenarios); active_s=s[s["Active"]].copy(); active_c=c[c["Active"]].copy(); sidx=active_s.set_index("ScopeID")
    batch=uuid.uuid4().hex; now=datetime.now(timezone.utc).isoformat(); solve_rows=[]; rating_rows=[]; worker_rows=[]; constraint_rows=[]
    pairs=active_c[["ScenarioID","ScopeID"]].drop_duplicates().sort_values(["ScenarioID","ScopeID"])
    for pair in pairs.itertuples(index=False):
        scope=sidx.loc[pair.ScopeID]; cfg=active_c[(active_c["ScenarioID"]==pair.ScenarioID)&(active_c["ScopeID"]==pair.ScopeID)].sort_values("PerformanceRating")
        pop=in_scope(w,scope); agg=aggregate(pop); hb=holdback(scope,sidx); gross=float(pop["EligibleVRTarget"].sum()); budget=gross*(1-hb); r=solve(agg,cfg,budget,spend_tolerance)
        meta=cfg.iloc[0]; recon="NOT_RUN"; delta=np.nan; hr=np.nan
        if r["status"]=="SUCCESS":
            ipm=r["ipm"]; hr=headroom(cfg,ipm)
            rr=agg.merge(cfg[["PerformanceRating","MinIPM","MaxIPM","ShapeScore"]],on="PerformanceRating"); rr["SolvedIPM"]=rr["PerformanceRating"].map(ipm); rr["RatingSpend"]=rr["EligibleVRTarget"]*rr["SolvedIPM"]
            for x in rr.itertuples():
                rating_rows.append(dict(BatchRunID=batch,ScenarioID=pair.ScenarioID,ScopeID=pair.ScopeID,RatingID=int(x.PerformanceRating),EligibleHC=int(x.EligibleHC),EligibleVRTarget=float(x.EligibleVRTarget),MinIPM=float(x.MinIPM),MaxIPM=float(x.MaxIPM),ShapeScore=float(x.ShapeScore),SolvedIPM=float(x.SolvedIPM),RatingSpend=float(x.RatingSpend),MinimumBindingFlag=abs(x.SolvedIPM-x.MinIPM)<=tol,MaximumBindingFlag=abs(x.MaxIPM-x.SolvedIPM)<=tol))
                for typ,val,actual,slack in [("Rating_Min",x.MinIPM,x.SolvedIPM,x.SolvedIPM-x.MinIPM),("Rating_Max",x.MaxIPM,x.SolvedIPM,x.MaxIPM-x.SolvedIPM)]:
                    constraint_rows.append(dict(BatchRunID=batch,ScenarioID=pair.ScenarioID,ScopeID=pair.ScopeID,ConstraintType=typ,RatingID=int(x.PerformanceRating),ConstraintValue=float(val),ActualValue=float(actual),Slack=float(slack),BindingFlag=abs(slack)<=tol,ViolationFlag=slack < -tol))
            constraint_rows.append(dict(BatchRunID=batch,ScenarioID=pair.ScenarioID,ScopeID=pair.ScopeID,ConstraintType="Budget_Max",RatingID=pd.NA,ConstraintValue=budget,ActualValue=r["spend"],Slack=budget-r["spend"],BindingFlag=abs(budget-r["spend"])<=tol,ViolationFlag=r["spend"]>budget+tol))
            wr=pop.copy(); wr["BatchRunID"]=batch; wr["ScenarioID"]=pair.ScenarioID; wr["ScopeID"]=pair.ScopeID; wr["RatingID"]=wr["PerformanceRating"]; wr["SolvedIPM"]=wr["PerformanceRating"].map(ipm); wr["CalculatedPayout"]=np.where(wr["Eligible"],wr["ProratedVRTarget"]*wr["SolvedIPM"],0.0); wr["ValidationStatus"]="PASS"
            cols=["BatchRunID","ScenarioID","ScopeID","WorkerID","RatingID","Eligible","ProratedVRTarget","SolvedIPM","CalculatedPayout","ValidationStatus"]
            if "ExclusionReason" in wr: cols.insert(6,"ExclusionReason")
            worker_rows.append(wr[cols]); delta=float(wr["CalculatedPayout"].sum()-r["spend"]); recon="PASS" if abs(delta)<=tol else "FAIL"
        solve_rows.append(dict(BatchRunID=batch,ScenarioID=pair.ScenarioID,ScopeID=pair.ScopeID,ScopeType=scope.ScopeType,GroupExecutive=scope.GroupExecutive,SubGroup=scope.SubGroup,ScenarioName=meta.ScenarioName,ShapeName=meta.ShapeName,DifferentiationLevel=meta.DifferentiationLevel,BaselineFlag=bool(meta.BaselineFlag),GrossBudget=gross,HoldbackPct=hb,HoldbackAmount=gross*hb,SolverBudget=budget,MinimumRequiredSpend=r.get("minimum",np.nan),FundingShortfall=r.get("shortfall",0.0),MaxFeasibleSpend=r.get("max_spend",0.0),SolvedSpend=r.get("spend",0.0),BudgetRemaining=budget-r.get("spend",0.0) if r["status"]=="SUCCESS" else np.nan,BudgetUtilisationPct=100*r.get("spend",0.0)/budget if budget>0 and r["status"]=="SUCCESS" else np.nan,EligibleHC=int(pop["Eligible"].sum()),ExcludedHC=int((~pop["Eligible"]).sum()),SolveStatus=r["status"],FeasibleFlag=r["status"]=="SUCCESS",CurveAnchor=r.get("anchor"),CurveScale=r.get("curve_scale"),PolicyFitScore=r.get("policy_fit"),ConstraintHeadroomScore=hr,ReconciliationStatus=recon,ReconciliationDelta=delta))
    fs=pd.DataFrame(solve_rows); fr=pd.DataFrame(rating_rows); fw=pd.concat(worker_rows,ignore_index=True) if worker_rows else pd.DataFrame(); fc=pd.DataFrame(constraint_rows)

    # Baseline comparisons and stability.
    base=active_c[active_c["BaselineFlag"]].drop_duplicates(["ScopeID","ScenarioID"]).set_index("ScopeID")["ScenarioID"].to_dict()
    fr["BaselineIPM"]=np.nan; fr["DeltaVsBaselineIPM"]=np.nan
    if not fw.empty:
        fw["BaselineIPM"]=np.nan; fw["BaselinePayout"]=np.nan; fw["PayoutDeltaVsBaseline"]=np.nan
    stability=[]
    for scope,bid in base.items():
        br=fr[(fr["ScopeID"]==scope)&(fr["ScenarioID"]==bid)][["RatingID","SolvedIPM"]].rename(columns={"SolvedIPM":"b"})
        for sid in fr.loc[fr["ScopeID"]==scope,"ScenarioID"].unique():
            m=fr[(fr["ScopeID"]==scope)&(fr["ScenarioID"]==sid)].merge(br,on="RatingID",how="left"); ix=fr[(fr["ScopeID"]==scope)&(fr["ScenarioID"]==sid)].index; fr.loc[ix,"BaselineIPM"]=m["b"].to_numpy(); fr.loc[ix,"DeltaVsBaselineIPM"]=m["SolvedIPM"].to_numpy()-m["b"].to_numpy(); rng=max(float(m["MaxIPM"].max()-m["MinIPM"].min()),1e-9); score=float(np.clip(100*(1-np.sqrt(np.mean((m["SolvedIPM"]-m["b"])**2))/rng),0,100)); stability.append(dict(BatchRunID=batch,ScenarioID=sid,ScopeID=scope,StabilityScore=score))
        if not fw.empty:
            bw=fw[(fw["ScopeID"]==scope)&(fw["ScenarioID"]==bid)][["WorkerID","SolvedIPM","CalculatedPayout"]].rename(columns={"SolvedIPM":"bipm","CalculatedPayout":"bp"})
            for sid in fw.loc[fw["ScopeID"]==scope,"ScenarioID"].unique():
                ix=fw[(fw["ScopeID"]==scope)&(fw["ScenarioID"]==sid)].index; m=fw.loc[ix,["WorkerID","CalculatedPayout"]].merge(bw,on="WorkerID",how="left"); fw.loc[ix,"BaselineIPM"]=m["bipm"].to_numpy(); fw.loc[ix,"BaselinePayout"]=m["bp"].to_numpy(); fw.loc[ix,"PayoutDeltaVsBaseline"]=m["CalculatedPayout"].to_numpy()-m["bp"].to_numpy()
    if stability: fs=fs.merge(pd.DataFrame(stability),on=["BatchRunID","ScenarioID","ScopeID"],how="left")
    else: fs["StabilityScore"]=np.nan
    bad=fs["SolveStatus"].eq("SUCCESS")&fs["ReconciliationStatus"].eq("FAIL"); fs.loc[bad,"SolveStatus"]="RECONCILIATION_ERROR"; fs.loc[bad,"FeasibleFlag"]=False

    # GE scenario roll-up. Any failed configured subgroup removes the GE scenario from Pareto consideration.
    ge=[]
    for (b,exe,sid),g in fs.groupby(["BatchRunID","GroupExecutive","ScenarioID"]):
        q=g[g["ScopeType"]=="SubGroup"]; q=q if not q.empty else g[g["ScopeType"]=="GE"]
        if q.empty or not (q["SolveStatus"].eq("SUCCESS").all() and q["ReconciliationStatus"].eq("PASS").all()): continue
        wt=q["SolverBudget"].to_numpy(float); wt=np.ones(len(q)) if wt.sum()<=0 else wt
        wav=lambda col: float(np.average(q[col].to_numpy(float),weights=wt)) if q[col].notna().all() else np.nan
        total_budget=float(q["SolverBudget"].sum()); total_spend=float(q["SolvedSpend"].sum())
        ge.append(dict(BatchRunID=b,GroupExecutive=exe,ScenarioID=sid,BudgetUtilisationPct=100*total_spend/total_budget if total_budget else 100,PolicyFitScore=wav("PolicyFitScore"),StabilityScore=wav("StabilityScore"),ConstraintHeadroomScore=wav("ConstraintHeadroomScore"),SolverBudget=total_budget,SolvedSpend=total_spend))
    gp=pareto(pd.DataFrame(ge)) if ge else pd.DataFrame()
    if not gp.empty: fs=fs.merge(gp[["BatchRunID","GroupExecutive","ScenarioID","ParetoRank","ParetoEfficientFlag","DominatedFlag","DominatedByScenarioID","ParetoMetricsUsed"]],on=["BatchRunID","GroupExecutive","ScenarioID"],how="left")

    ds=pd.DataFrame([dict(ScenarioID=sid,ScenarioName=g["ScenarioName"].iloc[0],ShapeSummary=(g["ShapeName"].iloc[0] if g["ShapeName"].nunique()==1 else "Mixed"),DifferentiationLevel=(g["DifferentiationLevel"].iloc[0] if g["DifferentiationLevel"].nunique()==1 else "Mixed"),BaselineFlag=bool(g["BaselineFlag"].any()),ActiveFlag=True) for sid,g in active_c.groupby("ScenarioID")])
    dscope=active_s[["ScopeID","ScopeType","GroupExecutive","SubGroup","ParentScopeID","Active"]].copy(); dr=pd.DataFrame({"RatingID":RATINGS,"RatingLabel":[f"Rating {r}" for r in RATINGS],"SortOrder":RATINGS})
    fact_cols={"PerformanceRating","ProratedVRTarget","Eligible","EligibleVRTarget","ExclusionReason"}; dw=w[[x for x in w.columns if x not in fact_cols]].drop_duplicates("WorkerID")
    audit=pd.DataFrame([dict(BatchRunID=batch,RunTimestamp=now,EngineVersion=VERSION,WorkerRowCount=len(w),ScopeCount=active_s["ScopeID"].nunique(),ScenarioCount=active_c["ScenarioID"].nunique(),ExpectedSolveCount=len(pairs),SuccessfulSolveCount=int(fs["SolveStatus"].eq("SUCCESS").sum()),InfeasibleSolveCount=int(fs["SolveStatus"].eq("INFEASIBLE_MINIMUM_SPEND").sum()),ErrorCount=int((~fs["SolveStatus"].isin(["SUCCESS","INFEASIBLE_MINIMUM_SPEND"])).sum()))])
    return {"DimScenario":ds,"DimScope":dscope,"DimRating":dr,"DimWorker":dw,"FactSolveResult":fs,"FactRatingResult":fr,"FactConstraintResult":fc,"FactWorkerResult":fw,"FactRunAudit":audit,"FactGEScenarioPareto":gp}


def write_outputs(out:Path,data):
    out.mkdir(parents=True,exist_ok=True)
    for name,df in data.items():
        if name=="FactWorkerResult":
            if not df.empty: df.to_parquet(out/name,index=False,compression="snappy",partition_cols=["BatchRunID","ScenarioID"])
        else: df.to_parquet(out/f"{name}.parquet",index=False,compression="snappy")


def main():
    p=argparse.ArgumentParser(); p.add_argument("--input",required=True,type=Path); p.add_argument("--output",required=True,type=Path); p.add_argument("--spend-tolerance",type=float,default=1.0); a=p.parse_args()
    if not a.input.exists(): print(f"Input not found: {a.input}",file=sys.stderr); return 2
    start=time.perf_counter()
    try:
        w=read_table(a.input,"tblWorkers"); s=read_table(a.input,"tblScopeConfig"); c=read_table(a.input,"tblScenarioConfig")
        data=run_engine(w,s,c,a.spend_tolerance); data["FactRunAudit"]["RuntimeSeconds"]=time.perf_counter()-start; write_outputs(a.output,data)
        sha=hashlib.sha256(a.input.read_bytes()).hexdigest(); manifest={"EngineVersion":VERSION,"Input":str(a.input),"InputSHA256":sha,"Rows":{k:len(v) for k,v in data.items()}}; (a.output/"run_manifest.json").write_text(json.dumps(manifest,indent=2))
        x=data["FactRunAudit"].iloc[0]; print(f"solves={x.ExpectedSolveCount} success={x.SuccessfulSolveCount} infeasible={x.InfeasibleSolveCount} errors={x.ErrorCount}"); return 0 if x.ErrorCount==0 else 1
    except Exception as e: print(f"ERROR: {e}",file=sys.stderr); return 1

if __name__=="__main__": raise SystemExit(main())
